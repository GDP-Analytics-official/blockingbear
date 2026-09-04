"""Ripristino dei grafici prodotti dalla sandbox (chat_anonymization:
_restore_chart, _swap_media).

Nel png le etichette di un grafico sono pixel: i segnaposto non sono testo e
non si sostituiscono. La via è la stessa del PDF rigenerato dal .docx — si
ripristina il SORGENTE (qui l'SVG gemello, dove le etichette sono testo) e si
RIDISEGNA il png sull'host. Il grafico incastonato in un documento si sistema
scambiando i byte dell'immagine, che sono gli stessi del png consegnato.

Niente OpenRouter e niente modello PII: il grafico e il documento li produce
la sandbox VERA con matplotlib, python-docx e soffice, esattamente come nel
turno di chat. Il registro è una mappa scritta a mano.

Cosa verifica:
  - l'SVG uscito da matplotlib con svg.fonttype="none" ha le etichette come
    testo (se un giorno non fosse più così, tutto il resto è inutile);
  - il png ripristinato disegna i VALORI REALI (il testo si rilegge dal
    vettoriale che il rasterizzatore riceve) e conserva le dimensioni;
  - i font di matplotlib diventano famiglie generiche, che MuPDF conosce;
  - il grafico incastonato nel .docx viene scambiato con quello ripristinato;
  - il PDF rigenerato dal .docx eredita il grafico giusto;
  - senza SVG sorgente il png resta com'è e viene dichiarato non
    ripristinabile (nessuna finta riuscita);
  - un grafico senza segnaposto non viene ridisegnato per niente.

Serve Docker attivo con l'immagine della sandbox:
    docker build -t blockingbear-sandbox:1 backend/sandbox
Uso:
    python backend/tests/chart_restore_test.py
"""

import hashlib
import io
import os
import shutil
import sys
import time
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parents[1]        # backend/
DATA = HERE / "data" / "test_chart"
os.environ["BLOCKINGBEAR_DATA_DIR"] = str(DATA)
# DB isolato quanto i file: mai la DATABASE_URL di esercizio (backend/.env).
# Vuota = sqlite nei dati del test; BLOCKINGBEAR_TEST_DATABASE_URL = usa-e-getta.
os.environ["DATABASE_URL"] = os.environ.get("BLOCKINGBEAR_TEST_DATABASE_URL", "")
os.environ["BLOCKINGBEAR_SANDBOX_POOL"] = "1"
shutil.rmtree(DATA, ignore_errors=True)
DATA.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(HERE))

import fitz                                                        # noqa: E402
from app import chat_anonymization as chat_anon                    # noqa: E402
from app.chat_anonymization import MediaPool, restore_artifact     # noqa: E402
from app.engine import convert                                     # noqa: E402
from app.engine.docx import extract_text as extract_docx           # noqa: E402
from app.openrouter import sandbox                                 # noqa: E402

# prefisso e profilo LibreOffice tutti nostri: il backend di sviluppo
# eventualmente attivo non deve accorgersi di questo test
sandbox._NAME_PREFIX = "blockingbear-sbxc-"
convert._PROFILE_DIR = DATA / "lo_profile"

MAPPING = {
    "[ORG_1]": "Fratelli Bianchi & Figli S.r.l.",
    "[ORG_2]": "Costruzioni Meridionali S.p.A.",
    "[FULLNAME_1]": "Mario Rossi",
    # molto più lungo del segnaposto: senza allargare la finestra dell'SVG
    # uscirebbe dal riquadro e il rasterizzatore lo taglierebbe a metà
    "[ORG_3]": "Consorzio Nazionale Costruzioni e Infrastrutture Integrate "
               "S.p.A.",
}

# Il codice che scrive il modello in una chat anonimizzata seguendo la regola
# di sistema (openrouter/chat.py, _ANON_PDF_RULE): figura salvata due volte,
# .svg e .png, e il .png incastonato nel documento.
CODE = '''
import matplotlib
matplotlib.use("Agg")
matplotlib.rcParams["svg.fonttype"] = "none"
import matplotlib.pyplot as plt

fig, ax = plt.subplots(figsize=(10, 5))
ax.bar(["[ORG_1]", "[ORG_2]"], [72, 45], color=["#3C6E9F", "#C9822B"])
ax.set_title("Fatturato per cliente - referente [FULLNAME_1]")
ax.set_ylabel("Migliaia di euro")
plt.tight_layout()
plt.savefig("/workspace/outputs/grafico.svg")
plt.savefig("/workspace/outputs/grafico.png", dpi=150)
plt.close()

# un secondo grafico SENZA segnaposto: non deve essere ridisegnato
fig, ax = plt.subplots(figsize=(4, 3))
ax.plot([1, 2, 3], [1, 4, 9])
ax.set_title("Andamento")
plt.savefig("/workspace/outputs/neutro.svg")
plt.savefig("/workspace/outputs/neutro.png", dpi=150)
plt.close()

# un terzo grafico a barre ORIZZONTALI con le etichette a sinistra, dove lo
# spazio è quello del segnaposto: è il caso che sbordava sul campo
fig, ax = plt.subplots(figsize=(8, 4))
ax.barh(["[ORG_3]", "[ORG_2]"], [120, 80])
ax.set_title("Fatturato per cliente")
plt.tight_layout()
plt.savefig("/workspace/outputs/lungo.svg")
plt.savefig("/workspace/outputs/lungo.png", dpi=150)
plt.close()

# un quarto grafico con le etichette su DUE RIGHE: matplotlib non usa più
# x/y + text-anchor, impagina lui riga per riga con transform="translate(...)"
fig, ax = plt.subplots(figsize=(9, 5))
ax.barh(["[ORG_3]\\n(marchio del bene ordinato)",
         "[ORG_2]\\n(cliente / intestatario fattura)",
         "[ORG_1]\\n(fornitore)"], [125, 152.5, 152.5])
ax.set_title("Soggetti coinvolti nell'ordine DO-2026-10701\\n(valori di test)")
ax.set_xlabel("EUR")
plt.tight_layout()
plt.savefig("/workspace/outputs/multiriga.svg")
plt.savefig("/workspace/outputs/multiriga.png", dpi=110)
plt.close()

from docx import Document
from docx.shared import Cm
import subprocess

doc = Document()
doc.add_heading("Riepilogo clienti", level=1)
doc.add_paragraph("Referente della pratica: [FULLNAME_1].")
doc.add_picture("/workspace/outputs/grafico.png", width=Cm(16))
t = doc.add_table(rows=2, cols=2)
t.cell(0, 0).text = "[ORG_1]"
t.cell(0, 1).text = "72"
t.cell(1, 0).text = "[ORG_2]"
t.cell(1, 1).text = "45"
doc.save("/workspace/outputs/riepilogo.docx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                "/workspace/outputs", "/workspace/outputs/riepilogo.docx"],
               check=True, capture_output=True, timeout=180)
print("ok")
'''

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}  {detail}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")


def drawn_text(svg, fit=True):
    """Il testo che il rasterizzatore disegnerà DAVVERO: MuPDF legge l'SVG
    come un documento e il suo layer di testo è esattamente ciò che finisce
    nei pixel — compreso il taglio di quel che esce dalla finestra, che è il
    difetto contro cui esiste `_fit_canvas`. Si passa dalle stesse due
    trasformazioni della rasterizzazione (`fit=False` ne salta una, per la
    controprova), o si misurerebbe un altro file."""
    prepared = chat_anon._svg_host_fonts(svg)
    if fit:
        prepared = chat_anon._fit_canvas(prepared)
    with fitz.open(stream=prepared.encode("utf-8"), filetype="svg") as doc:
        return "\n".join(page.get_text() for page in doc)


def png_size(data):
    pix = fitz.Pixmap(data)
    return pix.width, pix.height


def media(docx_bytes):
    """{nome: sha1} delle immagini incastonate in un OOXML."""
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        return {n: hashlib.sha1(z.read(n)).hexdigest()
                for n in z.namelist() if "/media/" in n.lower()}


def pdf_images(pdf_bytes):
    """[(byte, estensione)] delle immagini incastonate in un PDF."""
    out = []
    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        for page in doc:
            for img in page.get_images(full=True):
                info = doc.extract_image(img[0])
                out.append((info["image"], info["ext"]))
    return out


def distance(a, b, ext_a="png", ext_b="png", width=600):
    """Differenza media per canale tra due immagini, riportate alla stessa
    larghezza. LibreOffice RI-CODIFICA e RIDIMENSIONA le immagini quando
    converte in PDF (qui: png 1500px -> jpeg 1949px), quindi confrontare gli
    hash direbbe solo che i byte sono cambiati."""
    def thumb(data, ext):
        with fitz.open(stream=data, filetype=ext) as doc:
            page = doc[0]
            zoom = width / page.rect.width
            return page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)

    pa, pb = thumb(a, ext_a), thumb(b, ext_b)
    if (pa.width, pa.height) != (pb.width, pb.height):
        return 255.0
    return sum(abs(x - y) for x, y in zip(pa.samples, pb.samples)) / len(pa.samples)


def main():
    print("Avvio sandbox Docker...")
    sandbox.start()
    deadline = time.time() + 120
    while time.time() < deadline and not (sandbox.status()["available"]
                                          and sandbox.status()["pool_ready"]):
        time.sleep(0.5)
    status = sandbox.status()
    if not status["available"]:
        sys.exit(f"Docker non disponibile: {status}")

    print("Genero grafico e documento nella sandbox (matplotlib + "
          "python-docx + soffice)...")
    result = sandbox.execute("conv-chart-test", CODE, timeout=180)
    if result["outcome"] != "ok":
        sys.exit(f"Esecuzione fallita: {result['outcome']}\n"
                 f"{result['stderr'][:2000]}")
    work = DATA / "artifacts"
    work.mkdir(exist_ok=True)
    files = {}
    for rel, host in result["files"].items():
        dst = work / Path(rel).name
        shutil.copyfile(host, dst)
        files[dst.name] = dst
    print(f"  prodotti: {', '.join(sorted(files))} "
          f"in {result['elapsed_ms']} ms")
    for name in ("grafico.svg", "grafico.png", "riepilogo.docx",
                 "riepilogo.pdf"):
        if name not in files:
            sys.exit(f"La sandbox non ha prodotto {name}: {sorted(files)}")

    svg, png = files["grafico.svg"], files["grafico.png"]
    original_png = png.read_bytes()

    # 1. il presupposto: matplotlib con svg.fonttype="none" scrive TESTO
    raw_svg = svg.read_text("utf-8")
    check("nell'SVG di matplotlib le etichette sono testo, non curve",
          "<text" in raw_svg and "[ORG_1]" in drawn_text(raw_svg),
          f"{raw_svg.count('<text')} nodi <text>")

    # 2. il png ridisegnato dall'SVG sorgente
    data, n, left, extra = restore_artifact(png, MAPPING, source=svg)
    check("il grafico viene ridisegnato dall'SVG",
          data is not None and extra.get("chart") and n == 3,
          f"n={n} left={left} extra={extra}")
    if data is None:
        return summary()
    check("il png ripristinato è un png diverso dall'originale",
          data[:8] == b"\x89PNG\r\n\x1a\n" and data != original_png,
          f"{len(original_png)} -> {len(data)} byte")
    check("le dimensioni in pixel restano quelle del grafico originale",
          png_size(data) == png_size(original_png),
          f"{png_size(original_png)} -> {png_size(data)}")

    # 3. ciò che verrà disegnato è il valore reale (il testo si rilegge dal
    # vettoriale che il rasterizzatore riceve, prima che diventi pixel)
    restored_svg, n_svg = chat_anon._restore_string(
        raw_svg, chat_anon._mapping_for_syntax(".svg", MAPPING))
    drawn = drawn_text(restored_svg)
    missing = [MAPPING[ph] for ph in ("[ORG_1]", "[ORG_2]", "[FULLNAME_1]")
               if MAPPING[ph] not in drawn]
    check("le etichette disegnate sono i valori reali", not missing,
          f"mancanti: {missing}" if missing else drawn.split("\n")[0])
    check("nessun segnaposto resta nel grafico",
          not chat_anon._ANY_PH_RE.findall(drawn) and n_svg == 3,
          f"n={n_svg}")
    check("l'ampersand del valore non rompe l'XML dell'SVG",
          "&amp;" in restored_svg and "Fratelli Bianchi & Figli" in drawn)

    # 4. i font di matplotlib diventano famiglie che MuPDF conosce
    check("i font dell'SVG sono normalizzati per il rasterizzatore",
          "DejaVu" not in chat_anon._svg_host_fonts(restored_svg)
          and "sans-serif" in chat_anon._svg_host_fonts(restored_svg))
    check("l'SVG consegnato all'utente conserva i suoi font",
          "DejaVu" in restore_artifact(svg, MAPPING)[0].decode("utf-8"))

    # 4bis. l'etichetta molto più lunga del segnaposto non viene tagliata:
    # la finestra dell'SVG si allarga, il grafico resta delle stesse
    # dimensioni in pixel (o l'immagine scambiata nel documento si deformerebbe)
    long_svg = files["lungo.svg"].read_text("utf-8")
    long_png = files["lungo.png"].read_bytes()
    value = MAPPING["[ORG_3]"]
    restored_long = chat_anon._restore_string(
        long_svg, chat_anon._mapping_for_syntax(".svg", MAPPING))[0]
    long_data, long_n, _ll, _le = restore_artifact(
        files["lungo.png"], MAPPING, source=files["lungo.svg"])
    check("il caso serve davvero: senza allargare l'etichetta è tagliata",
          value not in drawn_text(restored_long, fit=False),
          drawn_text(restored_long, fit=False).replace("\n", " | ")[:120])
    check("con la finestra allargata l'etichetta lunga si legge intera",
          value in drawn_text(restored_long), f"n={long_n}")
    check("il grafico allargato conserva le dimensioni in pixel",
          long_data is not None and png_size(long_data) == png_size(long_png),
          f"{png_size(long_png)} -> {png_size(long_data or long_png)}")
    check("una finestra già capiente non viene toccata (idempotenza)",
          chat_anon._fit_canvas(chat_anon._fit_canvas(long_svg))
          == chat_anon._fit_canvas(long_svg))

    # 4ter. le etichette su più righe: matplotlib le posiziona riga per riga
    # con translate e senza ancora, cioè tenendo fermo il bordo SINISTRO. Se
    # non si riallinea, il valore cresce verso destra ed entra nel grafico
    # (etichette sopra le barre, tacca a mezzo del testo).
    multi_svg = files["multiriga.svg"].read_text("utf-8")

    def box_of(svg, needle):
        for attrs, text in chat_anon._SVG_TEXT_RE.findall(svg):
            if needle in chat_anon._svg_unescape(text):
                return chat_anon._text_box(attrs, text)
        return None

    def right_edge(svg, needle):
        return box_of(svg, needle)[2]

    blocks = chat_anon._text_blocks(multi_svg)
    label_blocks = [b for b in blocks if any("[ORG_" in n["text"] for n in b)]
    check("le etichette multiriga sono nodi con translate, non x/y",
          len(label_blocks) == 3 and all(len(b) == 2 for b in label_blocks),
          f"{len(blocks)} blocchi, righe {[len(b) for b in blocks]}")
    org2 = next(n for b in label_blocks for n in b if n["text"] == "[ORG_2]")
    check("il rettangolo di un nodo con translate è dove dice il translate",
          abs(box_of(multi_svg, "[ORG_2]")[0] - org2["x"]) < 0.01,
          f"misurato {box_of(multi_svg, '[ORG_2]')[0]:.1f}, "
          f"translate {org2['x']:.1f} (prima cadeva sull'origine)")
    check("l'allineamento si legge dalla geometria: etichette a destra",
          all(chat_anon._block_factor(b) == 1.0 for b in label_blocks),
          str([chat_anon._block_factor(b) for b in blocks]))
    check("il titolo su due righe resta centrato",
          any(chat_anon._block_factor(b) == 0.5
              for b in blocks if b not in label_blocks))

    plain = chat_anon._restore_string(
        multi_svg, chat_anon._mapping_for_syntax(".svg", MAPPING))[0]
    aligned = chat_anon._restore_svg(multi_svg, MAPPING)[0]
    was = right_edge(multi_svg, "[ORG_2]")
    check("il caso serve davvero: senza riallineare il bordo destro scappa",
          right_edge(plain, MAPPING["[ORG_2]"]) - was > 20,
          f"{was:.1f} -> {right_edge(plain, MAPPING['[ORG_2]']):.1f}")
    check("riallineata, l'etichetta conserva il bordo destro (non entra nel "
          "grafico)",
          abs(right_edge(aligned, MAPPING["[ORG_2]"]) - was) < 1.0,
          f"{was:.1f} -> {right_edge(aligned, MAPPING['[ORG_2]']):.1f}")
    check("le etichette multiriga si leggono per intero",
          all(MAPPING[ph] in drawn_text(aligned)
              for ph in ("[ORG_1]", "[ORG_2]", "[ORG_3]")),
          drawn_text(aligned).replace("\n", " | ")[:140])
    multi_out = restore_artifact(files["multiriga.png"], MAPPING,
                                 source=files["multiriga.svg"])[0]
    check("il grafico multiriga conserva le dimensioni in pixel",
          multi_out is not None
          and png_size(multi_out) == png_size(files["multiriga.png"].read_bytes()),
          f"{png_size(files['multiriga.png'].read_bytes())}")

    # 5. il grafico incastonato nel .docx
    charts = MediaPool()
    charts.add_pair(original_png, data)
    docx = files["riepilogo.docx"]
    before = media(docx.read_bytes())
    check("il .docx incastona esattamente il png consegnato",
          hashlib.sha1(original_png).hexdigest() in before.values(),
          str(before))
    doc_data, doc_n, doc_left, doc_extra = restore_artifact(
        docx, MAPPING, images=charts)
    check("il grafico dentro il .docx viene scambiato",
          doc_data is not None and doc_extra.get("images_swapped") == 1,
          str(doc_extra))
    after = media(doc_data) if doc_data else {}
    check("l'immagine scambiata è quella ripristinata",
          hashlib.sha1(data).hexdigest() in after.values(), str(after))
    doc_text = extract_docx(doc_data) if doc_data else ""
    check("anche il testo del .docx ha i valori reali",
          all(MAPPING[ph] in doc_text
              for ph in ("[ORG_1]", "[ORG_2]", "[FULLNAME_1]"))
          and not doc_left, f"n={doc_n} left={doc_left}")

    # 6. il PDF rigenerato dal .docx eredita il grafico giusto
    pdf = files["riepilogo.pdf"]
    pdf_data, _pn, _pl, pdf_extra = restore_artifact(
        pdf, MAPPING, source=docx, images=charts)
    check("il PDF viene rigenerato dal .docx sorgente",
          pdf_data is not None and pdf_extra.get("rerendered")
          and pdf_extra.get("images_swapped") == 1, str(pdf_extra))
    # Controprova a parità di conversione: lo STESSO documento senza lo
    # scambio, riconvertito da LibreOffice. Il png ripristinato deve essere
    # più vicino al grafico del PDF scambiato che a quello del PDF di
    # controllo — l'unica differenza tra i due è il testo delle etichette.
    control = convert.to_pdf(restore_artifact(docx, MAPPING)[0], suffix=".docx")
    inside = pdf_images(pdf_data or b"")
    ctrl = pdf_images(control)
    d_new = min((distance(img, data, ext) for img, ext in inside),
                default=255.0)
    d_old = min((distance(img, data, ext) for img, ext in ctrl), default=255.0)
    check("nel PDF c'è il grafico consegnato",
          len(inside) == 1 and d_new < 8.0, f"distanza={d_new:.2f}")
    check("è la versione ripristinata, non quella coi segnaposto",
          d_new < d_old,
          f"scambiato={d_new:.2f} controllo senza scambio={d_old:.2f}")

    # 7. i casi in cui NON si deve fingere di aver ripristinato
    check("senza SVG sorgente il png non è ripristinabile",
          restore_artifact(png, MAPPING) == (None, 0, [], {}))
    check("un sorgente che non è un SVG viene ignorato",
          restore_artifact(png, MAPPING, source=docx) == (None, 0, [], {}))
    neutral = files["neutro.png"]
    same, n0, _l0, extra0 = restore_artifact(neutral, MAPPING,
                                             source=files["neutro.svg"])
    check("un grafico senza segnaposto non viene ridisegnato",
          same == neutral.read_bytes() and n0 == 0 and not extra0)
    check("senza immagini da scambiare il .docx non dichiara scambi",
          not restore_artifact(docx, MAPPING)[3])

    keep = DATA / "grafico_ripristinato.png"
    keep.write_bytes(data)
    (DATA / "riepilogo_ripristinato.pdf").write_bytes(pdf_data or b"")
    print(f"\nGrafico ripristinato salvato per ispezione a mano: {keep}")
    return summary()


def summary():
    print(f"\n{PASS} PASS, {FAIL} FAIL")
    return 1 if FAIL else 0


if __name__ == "__main__":
    code = 1
    try:
        code = main()
    finally:
        sandbox.shutdown()
    sys.exit(code)
