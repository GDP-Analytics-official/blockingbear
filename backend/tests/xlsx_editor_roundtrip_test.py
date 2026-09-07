"""Cross-editor XLSX integration. Invoked by scripts/test-xlsx-roundtrip.sh.

openpyxl in the sandbox creates/edits the workbook; the backend redacts,
restores, and renders it with LibreOffice. No host packages or network calls.
"""
import sys
from pathlib import Path

MODE = sys.argv[1]
WORK = Path("/work")
MAPPING = {"[FULLNAME_1]": "O'Brien", "[FULLNAME_2]": "Mario Rossi",
           "[CUSTOM_1]": "PrivateTotal"}

if MODE in {"create", "edit", "check"}:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import PatternFill
    from openpyxl.workbook.defined_name import DefinedName

if MODE == "create":
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "O'Brien"
    ws["A1"] = "O'Brien"
    ws["A1"].fill = PatternFill("solid", fgColor="FF99CCFF")
    ws["A2"] = 7
    ws["B2"] = "='O''Brien'!A2+PrivateTotal"
    wb.defined_names.add(DefinedName("PrivateTotal", attr_text="'O''Brien'!$A$2"))
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=2, max_row=2))
    ws.add_chart(chart, "D1")
    ws["A3"] = "Jump"
    ws["A3"].hyperlink = "#'O''Brien'!A2"
    for i in range(7):
        other = wb.create_sheet(f"Sheet{i}")
        other["A1"] = i + 1
    ws = wb.create_sheet("Mario Rossi")
    ws["A1"] = "Mario Rossi"
    ws.sheet_state = "hidden"
    wb.save(WORK / "original.xlsx")
    print("PASS openpyxl created workbook with formulas, chart, links, hidden sheet and styles")

elif MODE == "redact":
    sys.path.insert(0, "/app")
    from app.engine.xlsx import redact_xlsx, known_xlsx_leaks, truncate_for_preview
    from app.engine.convert import to_pdf
    data, report = redact_xlsx((WORK / "original.xlsx").read_bytes(), MAPPING)
    assert not report["residual"], report
    assert report["frozen_formulas"] == 0, report
    assert not known_xlsx_leaks(data, MAPPING)
    (WORK / "protected.xlsx").write_bytes(data)
    pdf = to_pdf(truncate_for_preview(data)[0], suffix=".xlsx")
    assert pdf.startswith(b"%PDF")
    (WORK / "protected.pdf").write_bytes(pdf)
    print("PASS backend redaction and LibreOffice rendering")

elif MODE == "edit":
    wb = openpyxl.load_workbook(WORK / "protected.xlsx")
    assert len(wb.sheetnames) == 9
    assert wb.sheetnames[0].startswith("BB_S_")
    assert wb.sheetnames[8].startswith("BB_S_")
    assert wb.worksheets[8].sheet_state == "hidden"
    ws = wb.worksheets[0]
    assert ws["A1"].value == "[FULLNAME_1]"
    assert "BB_S_" in ws["B2"].value and "BB_N_" in ws["B2"].value
    assert "BB_S_" in ws._charts[0].series[0].val.numRef.f
    assert "BB_S_" in ws["A3"].hyperlink.target
    assert any(n.startswith("_BB_RESTORE_") for n in wb.defined_names)
    ws["C2"] = "Edited by the sandbox"
    wb.save(WORK / "edited.xlsx")
    print("PASS openpyxl loaded, checked and saved protected workbook")

elif MODE == "restore":
    sys.path.insert(0, "/app")
    from app.chat_anonymization import restore_artifact
    from app.engine.xlsx import known_xlsx_leaks
    from app.engine.convert import to_pdf
    assert not known_xlsx_leaks((WORK / "edited.xlsx").read_bytes(), MAPPING)
    data, n, left, extra = restore_artifact(WORK / "edited.xlsx", MAPPING)
    assert data is not None and n > 0 and not left, (n, left, extra)
    (WORK / "restored.xlsx").write_bytes(data)
    pdf = to_pdf(data, suffix=".xlsx")
    assert pdf.startswith(b"%PDF")
    print("PASS backend restoration after editor save and LibreOffice rendering")

elif MODE == "check":
    wb = openpyxl.load_workbook(WORK / "restored.xlsx")
    original = openpyxl.load_workbook(WORK / "original.xlsx")
    assert wb.sheetnames == original.sheetnames
    assert wb.worksheets[8].sheet_state == "hidden"
    ws = wb.worksheets[0]
    assert ws["A1"].value == "O'Brien"
    fill, expected = ws["A1"].fill, original.worksheets[0]["A1"].fill
    # StyleProxy identity and an unused solid-fill background are not visual
    # formatting. The redactor stores its reversible marker in that background.
    assert fill.patternType == expected.patternType == "solid"
    assert fill.fgColor == expected.fgColor, (fill.fgColor, expected.fgColor)
    assert ws["B2"].value == original.worksheets[0]["B2"].value
    assert ws["C2"].value == "Edited by the sandbox"
    assert ws._charts[0].series[0].val.numRef.f == original.worksheets[0]._charts[0].series[0].val.numRef.f
    assert ws["A3"].hyperlink.target == original.worksheets[0]["A3"].hyperlink.target
    assert wb.defined_names["PrivateTotal"].attr_text == original.defined_names["PrivateTotal"].attr_text
    assert not any(n.startswith("_BB_RESTORE_") for n in wb.defined_names)
    print("PASS restored names, formulas, chart, links, edits, styles and hidden state")
else:
    raise SystemExit("Unknown test phase")
