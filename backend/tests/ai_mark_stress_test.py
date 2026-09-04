"""Stress test della marcatura AI Act (engine/ai_mark.py) su file generati
DAVVERO dalla sandbox Docker, con gli stessi produttori che usa il modello:

  docx  python-docx x3 (semplice / tabella+stili / immagine incorporata)
        + LibreOffice (roundtrip odt -> docx)
  xlsx  openpyxl (multi-foglio, formule, celle unite) / xlsxwriter (grafico
        incorporato) / pandas.to_excel + LibreOffice (ods -> xlsx)
  pptx  python-pptx x3 (titolo / tabella+forme / immagine)
        + LibreOffice (odp -> pptx)
  pdf   soffice da docx / matplotlib savefig / pypdf (fusione)
  png   matplotlib x2 (dpi diversi) / pillow RGBA / pillow palette
  svg   matplotlib x2 (fonttype none e default) / scritto a mano
  odf   odt/ods/odp: NON marcabili per scelta — si pretende file intatto

Per ogni file: marcatura riuscita, contenuto INTATTO (part OOXML byte a
byte, testo+pagine PDF, pixel PNG), idempotenza, e riconversione LibreOffice
dei marcati. In coda una batteria di concorrenza: 8 thread marcano in
parallelo mentre un altro renderizza pagine con fitz (il lock MuPDF regge).

Serve Docker attivo e l'immagine blockingbear-sandbox:1.
Uso:
    python backend/tests/ai_mark_stress_test.py
"""
import os
import shutil
import sys
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from xml.etree import ElementTree as ET

HERE = Path(__file__).resolve().parents[1]        # backend/
sys.path.insert(0, str(HERE))

os.environ["BLOCKINGBEAR_DATA_DIR"] = str(HERE / "data" / "test_aimark_stress")
# DB isolato quanto i file: mai la DATABASE_URL di esercizio (backend/.env).
# Vuota = sqlite nei dati del test; BLOCKINGBEAR_TEST_DATABASE_URL = usa-e-getta.
os.environ["DATABASE_URL"] = os.environ.get("BLOCKINGBEAR_TEST_DATABASE_URL", "")
os.environ["BLOCKINGBEAR_SANDBOX_POOL"] = "1"

from app.db import init_db                                        # noqa: E402
from app.engine import ai_mark, convert                           # noqa: E402
from app.openrouter import sandbox                                # noqa: E402

import fitz                                                       # noqa: E402
from PIL import Image                                             # noqa: E402

sandbox._NAME_PREFIX = "blockingbear-sbxs-"     # mai toccare il pool di esercizio

WORK = HERE / "data" / "test_aimark_stress" / "files"
CONV = "stress"

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}  {detail}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")


def run(code, timeout=120):
    """Esegue nella sandbox e copia subito gli artifact in WORK (lo staging
    muore col container). Torna (outcome, [nomi copiati])."""
    r = sandbox.execute(CONV, code, timeout=timeout)
    copied = []
    for rel, host in r.get("files", {}).items():
        dst = WORK / Path(rel).name
        shutil.copyfile(host, dst)
        copied.append(dst.name)
    if r["outcome"] != "ok":
        print(f"    [sandbox] {r['outcome']}: {r.get('stderr', '')[:400]}")
    return r["outcome"], copied


# --- Generazione nella sandbox --------------------------------------------------

GEN_STEPS = [
    ("python-docx x3", """
from docx import Document
from docx.shared import Pt, Inches
from PIL import Image

d = Document()
d.add_heading('Relazione però semplice', 0)
for i in range(30):
    d.add_paragraph(f'Paragrafo {i} con çàràttèri e €urosimboli ✓.')
d.save('/workspace/outputs/docx_semplice però.docx')

d = Document()
d.add_heading('Tabella e stili', 1)
t = d.add_table(rows=12, cols=5)
t.style = 'Light Grid Accent 1'
for r_i, row in enumerate(t.rows):
    for c_i, cell in enumerate(row.cells):
        cell.text = f'r{r_i}c{c_i}'
p = d.add_paragraph()
run = p.add_run('Grassetto enorme')
run.bold = True
run.font.size = Pt(28)
d.save('/workspace/outputs/docx_tabella.docx')

Image.new('RGB', (300, 120), (200, 40, 40)).save('/tmp/img.png')
d = Document()
d.add_paragraph('Documento con immagine incorporata:')
d.add_picture('/tmp/img.png', width=Inches(3))
d.add_page_break()
d.add_paragraph('Seconda pagina.')
d.save('/workspace/outputs/docx_immagine.docx')
print('ok')
"""),
    ("openpyxl / xlsxwriter / pandas", """
import openpyxl, xlsxwriter
import pandas as pd, numpy as np

wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Dati'
for r in range(1, 101):
    for c in range(1, 9):
        ws.cell(row=r, column=c, value=r * c)
ws['J1'] = '=SUM(A1:H100)'
ws.merge_cells('A105:H106')
ws['A105'] = 'celle unite però lünghe'
wb.create_sheet('Vuoto')
wb.save('/workspace/outputs/xlsx_openpyxl.xlsx')

wb = xlsxwriter.Workbook('/workspace/outputs/xlsx_xlsxwriter.xlsx')
ws = wb.add_worksheet('Grafico')
bold = wb.add_format({'bold': True, 'bg_color': '#DDEBF7'})
ws.write_row(0, 0, ['mese', 'valore'], bold)
for i, v in enumerate([5, 9, 3, 7, 12]):
    ws.write_row(i + 1, 0, [f'm{i}', v])
ch = wb.add_chart({'type': 'column'})
ch.add_series({'values': '=Grafico!$B$2:$B$6'})
ws.insert_chart('D2', ch)
wb.close()

df = pd.DataFrame(np.random.default_rng(7).integers(0, 1000, (200, 12)),
                  columns=[f'col_{i}' for i in range(12)])
df.to_excel('/workspace/outputs/xlsx_pandas.xlsx', index=False)
print('ok')
"""),
    ("python-pptx x3", """
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = 'Titolo però àccentato'
s.placeholders[1].text = 'sottotitolo'
p.save('/workspace/outputs/pptx_titolo.pptx')

p = Presentation()
s = p.slides.add_slide(p.slide_layouts[6])
tb = s.shapes.add_table(6, 4, Inches(1), Inches(1),
                        Inches(7), Inches(3)).table
for r in range(6):
    for c in range(4):
        tb.cell(r, c).text = f'{r},{c}'
s.shapes.add_shape(1, Inches(1), Inches(5), Inches(3), Inches(1))
p.save('/workspace/outputs/pptx_tabella.pptx')

Image.new('RGB', (640, 360), (30, 60, 120)).save('/tmp/slide.png')
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[6])
s.shapes.add_picture('/tmp/slide.png', Inches(1), Inches(1), Inches(6))
p.save('/workspace/outputs/pptx_immagine.pptx')
print('ok')
"""),
    ("matplotlib png/svg/pdf + pillow", """
import matplotlib
matplotlib.use('Agg')
matplotlib.rcParams['svg.fonttype'] = 'none'
import matplotlib.pyplot as plt
from PIL import Image

fig, ax = plt.subplots(figsize=(7, 4))
ax.plot([1, 3, 2, 5], label='sèrie però lünga')
ax.legend(); ax.set_title('Grafico €')
fig.savefig('/workspace/outputs/png_mpl_line.png')
fig.savefig('/workspace/outputs/svg_fonttype.svg')
fig.savefig('/workspace/outputs/pdf_matplotlib.pdf')
plt.close(fig)

matplotlib.rcParams['svg.fonttype'] = 'path'
fig, ax = plt.subplots(figsize=(8, 5))
ax.barh(['alfa', 'beta', 'gamma'], [3, 1, 4])
fig.savefig('/workspace/outputs/png_mpl_barh.png', dpi=150)
fig.savefig('/workspace/outputs/svg_default.svg')
plt.close(fig)

Image.new('RGBA', (256, 128), (10, 200, 100, 128)).save(
    '/workspace/outputs/png_pil_rgba.png')
Image.new('RGB', (64, 64), (255, 0, 0)).convert(
    'P', palette=Image.ADAPTIVE).save('/workspace/outputs/png_pil_palette.png')

svg = ('<?xml version="1.0" encoding="utf-8"?>\\n'
       '<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">'
       '<rect width="10" height="10" fill="teal"/></svg>')
open('/workspace/outputs/svg_mano.svg', 'w').write(svg)
print('ok')
"""),
    ("soffice: pdf + odf + roundtrip OOXML", """
import shutil, subprocess

def conv(target, src):
    r = subprocess.run(['soffice', '--headless', '--convert-to', target,
                        '--outdir', '/tmp/conv', src],
                       capture_output=True, text=True, timeout=150)
    if r.returncode != 0:
        raise RuntimeError(r.stderr[-300:])

conv('pdf', '/workspace/outputs/docx_tabella.docx')
shutil.copy('/tmp/conv/docx_tabella.pdf', '/workspace/outputs/pdf_soffice.pdf')

conv('odt', '/workspace/outputs/docx_semplice però.docx')
conv('ods', '/workspace/outputs/xlsx_openpyxl.xlsx')
conv('odp', '/workspace/outputs/pptx_titolo.pptx')
shutil.copy('/tmp/conv/docx_semplice però.odt', '/workspace/outputs/odf_doc.odt')
shutil.copy('/tmp/conv/xlsx_openpyxl.ods', '/workspace/outputs/odf_calc.ods')
shutil.copy('/tmp/conv/pptx_titolo.odp', '/workspace/outputs/odf_pres.odp')

conv('docx', '/workspace/outputs/odf_doc.odt')
conv('xlsx', '/workspace/outputs/odf_calc.ods')
conv('pptx', '/workspace/outputs/odf_pres.odp')
shutil.copy('/tmp/conv/odf_doc.docx', '/workspace/outputs/docx_libreoffice.docx')
shutil.copy('/tmp/conv/odf_calc.xlsx', '/workspace/outputs/xlsx_libreoffice.xlsx')
shutil.copy('/tmp/conv/odf_pres.pptx', '/workspace/outputs/pptx_libreoffice.pptx')
print('ok')
""", 240),
    ("pypdf: fusione", """
from pypdf import PdfReader, PdfWriter
w = PdfWriter()
for src in ('/workspace/outputs/pdf_soffice.pdf',
            '/workspace/outputs/pdf_matplotlib.pdf'):
    for page in PdfReader(src).pages:
        w.add_page(page)
w.add_metadata({'/Title': 'Unione però'})
with open('/workspace/outputs/pdf_pypdf.pdf', 'wb') as f:
    w.write(f)
print('ok')
"""),
]


# --- Verifiche host --------------------------------------------------------------

def zip_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def verify_ooxml(path):
    before = zip_parts(path)
    check(f"{path.name}: marcato", ai_mark.mark_file(path) is True)
    after = zip_parts(path)
    core = after.get("docProps/core.xml", b"")
    check(f"{path.name}: dicitura in core.xml",
          ai_mark.MARK_TEXT.encode() in core)
    same = (list(before) == list(after)
            and all(before[n] == after[n] for n in before
                    if n != "docProps/core.xml"))
    check(f"{path.name}: altre part byte-identiche", same)
    try:
        ET.fromstring(core)
        ok_xml = True
    except ET.ParseError:
        ok_xml = False
    check(f"{path.name}: core.xml valido", ok_xml)
    size = path.stat().st_size
    check(f"{path.name}: idempotente",
          ai_mark.mark_file(path) is True and path.stat().st_size == size)


def verify_pdf(path):
    original = path.read_bytes()
    with fitz.open(str(path)) as doc:
        pages = doc.page_count
        text = [doc[i].get_text() for i in range(pages)]
    check(f"{path.name}: marcato", ai_mark.mark_file(path) is True)
    check(f"{path.name}: byte originali intatti (prefisso)",
          path.read_bytes().startswith(original))
    with fitz.open(str(path)) as doc:
        check(f"{path.name}: Subject+Keywords",
              ai_mark.MARK_TEXT in (doc.metadata.get("subject") or "")
              and "trainedAlgorithmicMedia" in (doc.metadata.get("keywords") or ""))
        check(f"{path.name}: pagine e testo identici",
              doc.page_count == pages
              and [doc[i].get_text() for i in range(pages)] == text)
    size = path.stat().st_size
    check(f"{path.name}: idempotente",
          ai_mark.mark_file(path) is True and path.stat().st_size == size)


def verify_png(path):
    with Image.open(path) as im:
        pixels = im.tobytes()
        mode = im.mode
    check(f"{path.name}: marcato", ai_mark.mark_file(path) is True)
    with Image.open(path) as im:
        im.load()
        check(f"{path.name}: chunk tEXt",
              im.info.get("AI-Generated") == ai_mark.MARK_TEXT)
        check(f"{path.name}: pixel identici ({mode})",
              im.mode == mode and im.tobytes() == pixels)
    size = path.stat().st_size
    check(f"{path.name}: idempotente",
          ai_mark.mark_file(path) is True and path.stat().st_size == size)


def verify_svg(path):
    check(f"{path.name}: marcato", ai_mark.mark_file(path) is True)
    text = path.read_text("utf-8")
    try:
        root = ET.fromstring(text)
        ok = root.tag.endswith("svg")
    except ET.ParseError:
        ok = False
    check(f"{path.name}: XML valido con dicitura",
          ok and ai_mark.MARK_TEXT in text)
    check(f"{path.name}: idempotente",
          ai_mark.mark_file(path) is True
          and path.read_text("utf-8") == text)


def verify_odf(path):
    before = path.read_bytes()
    check(f"{path.name}: NON marcabile e intatto",
          ai_mark.mark_file(path) is False and path.read_bytes() == before)


def verify_lo_reconvert():
    """LibreOffice (host) riapre e converte OGNI OOXML marcato."""
    print("\n--- LibreOffice riconverte i marcati ---")
    if not convert.available():
        print("  SKIP  LibreOffice non trovato sul host")
        return
    for path in sorted(WORK.glob("*.docx")) + sorted(WORK.glob("*.xlsx")) \
            + sorted(WORK.glob("*.pptx")):
        pdf = None
        for attempt in (1, 2):
            try:
                pdf = convert.to_pdf(path.read_bytes(), suffix=path.suffix)
                break
            except convert.ConvertError as e:
                if attempt == 2:
                    check(f"{path.name} -> PDF", False, repr(e))
        if pdf is not None:
            check(f"{path.name} -> PDF", pdf.startswith(b"%PDF"),
                  f"{len(pdf)} byte")


def stress_concurrency():
    """8 thread marcano copie fresche in parallelo mentre un altro thread
    renderizza pagine con fitz: il lock MuPDF deve reggere (il crash nativo
    di PyMuPDF senza lock è il guasto storico, vedi mupdf_lock.py)."""
    print("\n--- Concorrenza (lock MuPDF) ---")
    pool = [p for p in WORK.iterdir()
            if p.suffix in (".pdf", ".docx", ".xlsx", ".pptx", ".png", ".svg")]
    arena = WORK / "arena"
    arena.mkdir(exist_ok=True)
    copies = []
    for i in range(48):
        src = pool[i % len(pool)]
        dst = arena / f"{i:02d}_{src.name}"
        shutil.copyfile(src, dst)
        copies.append(dst)

    stop = time.monotonic() + 60
    renders = [0]

    def hammer():
        sample = next(p for p in pool if p.suffix == ".pdf")
        while time.monotonic() < stop and renders[0] >= 0:
            with fitz.open(str(sample)) as doc:
                doc[0].get_pixmap(dpi=60)
            renders[0] += 1

    with ThreadPoolExecutor(max_workers=9) as ex:
        h = ex.submit(hammer)
        results = list(ex.map(ai_mark.mark_file, copies))
        renders[0] = -abs(renders[0]) - 1        # ferma il martello
        h.result()

    check("48 marcature concorrenti riuscite", all(results),
          f"{sum(bool(r) for r in results)}/48, render fitz in parallelo: "
          f"{-renders[0] - 1}")
    # spot check d'integrità su un campione di ogni tipo
    for ext, fn in ((".pdf", None), (".docx", None), (".png", None)):
        p = next(c for c in copies if c.suffix == ext)
        if ext == ".pdf":
            with fitz.open(str(p)) as doc:
                ok = ai_mark.MARK_TEXT in (doc.metadata.get("subject") or "")
        elif ext == ".docx":
            ok = ai_mark.MARK_TEXT.encode() in zip_parts(p)["docProps/core.xml"]
        else:
            with Image.open(p) as im:
                im.load()
                ok = im.info.get("AI-Generated") == ai_mark.MARK_TEXT
        check(f"concorrenza: {p.name[3:]} integro e marcato", ok)


def main():
    root = HERE / "data" / "test_aimark_stress"
    if root.is_dir():
        shutil.rmtree(root, ignore_errors=True)
    WORK.mkdir(parents=True)
    init_db()

    print("--- Generazione nella sandbox Docker ---")
    sandbox.start()
    deadline = time.time() + 60
    while time.time() < deadline and sandbox.status()["detecting"]:
        time.sleep(0.5)
    if not sandbox.status()["available"]:
        print("Docker non disponibile: impossibile generare i file. STOP.")
        return 1
    try:
        for entry in GEN_STEPS:
            name, code = entry[0], entry[1]
            timeout = entry[2] if len(entry) > 2 else 120
            outcome, copied = run(code, timeout=timeout)
            check(f"generazione {name}", outcome == "ok",
                  f"{len(copied)} file")
    finally:
        sandbox.shutdown()

    by_ext = {}
    for p in sorted(WORK.iterdir()):
        by_ext.setdefault(p.suffix.lower(), []).append(p)
    counts = {e: len(v) for e, v in sorted(by_ext.items())}
    print(f"\nGenerati: {counts}")
    for ext, minimo in ((".docx", 4), (".xlsx", 4), (".pptx", 4),
                        (".pdf", 3), (".png", 4), (".svg", 3)):
        check(f"almeno {minimo} file {ext}", len(by_ext.get(ext, [])) >= minimo)

    print("\n--- OOXML ---")
    for p in by_ext.get(".docx", []) + by_ext.get(".xlsx", []) \
            + by_ext.get(".pptx", []):
        verify_ooxml(p)
    print("\n--- PDF ---")
    for p in by_ext.get(".pdf", []):
        verify_pdf(p)
    print("\n--- PNG ---")
    for p in by_ext.get(".png", []):
        verify_png(p)
    print("\n--- SVG ---")
    for p in by_ext.get(".svg", []):
        verify_svg(p)
    print("\n--- ODF (non marcabili, intatti) ---")
    for ext in (".odt", ".ods", ".odp"):
        for p in by_ext.get(ext, []):
            verify_odf(p)

    verify_lo_reconvert()
    stress_concurrency()

    print(f"\n{PASS} PASS, {FAIL} FAIL")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
